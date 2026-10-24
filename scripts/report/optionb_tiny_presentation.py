"""Build a2-slide epoch44 presentation from bound benchmark evidence.

Requires Python 3.12, python-pptx, Pillow and reportlab.
Uses the existing genuine Manim clips; never imports a training/runtime model.
The PDF is an image-based visual companion, not a PowerPoint-to-PDF conversion.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import statistics
import subprocess
import zipfile
from pathlib import Path

MODEL_SHA = "c1cd773b92337fd2a453554b047fa555c692fb981b5975700a32260388d4b9ce"
SOURCE_COMMIT = "ad4d6262e7f91b2cb4cb4f75a1b1b082d6b088f3"
TAG = "optionb-tiny-v1.0.0-rc.1"
PREFIX = "optionb-tiny-epoch44-presentation"
REPO = "justinruaya123/smart-speaker"
ROOT = Path(__file__).resolve().parents[2]
BG, INK, MUTED, TEAL, AMBER = "#090D11", "#E2ECEF", "#96ADB7", "#9CE8D3", "#E8B36E"

SCRIPTS = [
    "Wake detection belongs inside inference. OpenWakeWord is an experimental wake adapter here, not installed or promoted because its regression failed. The native wake backend stays unchanged.\n\n"
    "After a wake, VAD captures one command. An independent speaker and overlap check guards ownership. The command model is the original epoch 44 Tiny graph: 40 log-mel bins, frozen normalization, stride two, 160 channels and four causal residual blocks with dilations one, two, four and eight.\n\n"
    "Masked mean and max pooling feed 32 logits: 31 complete intent-and-slot tuples plus UNKNOWN. The probability and margin thresholds are 0.50 and 0.10. A brightness 60 class already fixes light.all and percent 60: no transcript or slot decoder. Ownership and typed validation still guard any action. The animation is a schematic, not a live recognition trace.",
    "These figures measure the frozen command stage, not OpenWakeWord, VAD, ownership, UI or actions. Of 4,614 replayed originals, 187 ambiguous references stay unscored; timing repetitions do not add accuracy samples.\n\n"
    "Correct accepted recall is 64.53 percent on human TEST and 90 percent on the smaller, exposed HOLDOUT. TEST negatives produce 24 false accepts out of 256. Synthetic TEST recall is 97.65 percent, not proof of human reliability.\n\n"
    "On an i7 CPU, the command stage takes six milliseconds median and under ten at the 95th percentile. The isolated process peaks at 73.61 mebibytes.\n\n"
    "This is a research prerelease. Wake qualification failed; fresh blind-human, live-microphone and Raspberry Pi qualification remain future work. Default actions stay off.",
]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def prepare(args):
    bench = args.benchmark
    result = read_json(bench / "results.json")
    resources = read_json(bench / "resource-results.json")
    frozen = args.frozen_source or ROOT
    release_protocol = frozen / "experiments/speaker-optionb-release-v1/protocol-v3.json"
    if not release_protocol.is_file():
        raise FileNotFoundError("Pass --frozen-source pointing to the pinned original source checkout; do not replace Iris main modules.")
    architecture = read_json(release_protocol)["architecture"]
    training = read_json(frozen / "experiments/speaker-optionb-tiny-v1/training-summary.json")
    model = result["model"]
    assert model["bindings"]["model.onnx"] == MODEL_SHA == resources["graph_sha256"]
    assert sha(ROOT / "submission/models/optionb-tiny/model.onnx") == MODEL_SHA
    assert model["epoch"] ==44 and model["parameters"] ==258912
    assert model["policy"]["min_probability"] ==.5 and model["policy"]["min_margin"] ==.1
    assert result["all_original_decisions_and_repeats_match_frozen"]
    assert result["max_original_logit_abs_difference"] ==0
    assert not result["actions_enabled"] and result["exposed_regression_not_blind"]
    observations = [json.loads(line) for line in (bench / "observations.jsonl").read_text().splitlines()]
    assert len(observations) == result["timing_observations"] ==3*result["original_recordings"]
    human = [r["latency_ms"] for r in observations if r["kind"] == "human"]
    human_p95 = statistics.quantiles(human, n=100, method="inclusive")[94]
    expected = {"test/human": (393,609,24,256), "holdout/human": (72,80,1,14),
                "test/synthetic": (3289,3368,0,0), "holdout/synthetic": (99,100,0,0)}
    for cohort, (tp, positives, fa, negatives) in expected.items():
        s = result["cohorts"][cohort]["selected"]
        assert (s["tp"],s["positives"],s["negative_false_accepts"],s["negatives"]) == (tp,positives,fa,negatives)
    commands = read_json(ROOT / "submission/models/optionb-tiny/commands.json")
    example = {"intent": "BRIGHTNESS", "slots": {"light_id": "light.all", "percent":60}}
    assert len(commands) ==32 and commands[0] is None and example in commands
    findings_path = ROOT / "experiments/iris-openwakeword-publication-v1/findings.json"
    findings = read_json(findings_path)
    assert findings["decision"] == "DO_NOT_PROMOTE"
    assert findings["runtime"]["integration_performed"] is False
    assert not findings["boundary"]["installed"] and not findings["boundary"]["promoted"]
    scored = sum(c["selected"]["rows"] for c in result["cohorts"].values())
    ambiguous = sum(c["ambiguous_unscored"] for c in result["cohorts"].values())
    assert scored + ambiguous == result["original_recordings"]
    data = {
        "format": "optionb-tiny-2minute-presentation-v1", "repository": REPO,
        "release_tag": TAG, "model_source_commit": SOURCE_COMMIT, "model": model,
        "architecture": architecture, "illustrative_command": example,
        "training": {k: training[k] for k in ("selected_epoch","selected_on","epochs_completed","training_config")},
        "cohorts": result["cohorts"], "original_recordings": result["original_recordings"],
        "scored_original_recordings":scored, "ambiguous_original_recordings":ambiguous,
        "timing_observations": len(observations), "hardware": result["hardware"],
        "human_latency_ms": {"n":len(human),"p50":statistics.median(human),"p95":human_p95},
        "all_recordings_latency_ms": result["latency_ms"], "resource_probe": resources,
        "timing_excludes": result["timing_excludes"], "exposed_regression_not_blind": True,
        "actions_enabled": False, "fresh_blind_human_pi_or_microphone_qualification": False,
        "source_sha256": {name:sha(bench/name) for name in ("results.json","resource-results.json","protocol.json","observations.jsonl")},
        "architecture_protocol_sha256":sha(release_protocol),
        "inference": {
            "publication_branch":"main",
            "openwakeword": {
                "role":"wake_detection", "runtime_module":"scripts/runtime/iris_openwakeword.py",
                "integration_performed":False,"promoted":False,"qualified":False,
                "findings_sha256":sha(findings_path),
                "status":"Rejected experimental alternative; existing native wake remains unchanged",
            },
            "flow":["microphone","wake detection","command VAD","command model and independent ownership","typed validation","configured action and receipt"],
        },
        "benchmark_scope": {
            "stage":"frozen_epoch44_command_only",
            "excludes":["wake","VAD/capture","owner","UI","providers","speaking","file I/O"],
        },
        "architecture_media_sha256": {name:sha(ROOT/name) for name in (
            "presentation/iris_scenes.py","docs/site/animations/pipeline.mp4","docs/site/animations/model.mp4")},
        "slides": [{"number":i+1,"target_seconds":s,"script":text,
                    "script_word_count":len(text.split())}
                   for i,(s,text) in enumerate(zip((65,55),SCRIPTS))],
        "suggested_total_seconds":120,
        "timing_note":"A suggested speaking budget, not a measured human delivery or voice recording.",
    }
    assert sum(s["target_seconds"] for s in data["slides"]) ==120
    assert len(data["slides"]) ==2
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(data,indent=2)+"\n",encoding="utf-8")
    documentation = ROOT / "docs/PRESENTATION-optionb-tiny-2minute.md"
    if documentation.is_file():
        text = documentation.read_text(encoding="utf-8")
        narration = "## Narration\n\n" + "\n\n".join(
            f"### Slide {s['number']} — target {s['target_seconds']} seconds\n\n{s['script']}"
            for s in data["slides"]
        ) + "\n\n"
        text = re.sub(r"## Narration\n.*?(?=## Source and rebuild)",lambda _:narration,text,flags=re.S)
        counts = [s["script_word_count"] for s in data["slides"]]
        text = re.sub(r"Speaker notes contain .*?for slide 2\.",
                      f"Speaker notes contain {counts[0]} words for slide 1 and {counts[1]} for slide 2.",text)
        documentation.write_text(text,encoding="utf-8",newline="\n")
    print(json.dumps({"data":str(args.output),"human_latency_ms":data["human_latency_ms"],
                      "script_words":[s["script_word_count"] for s in data["slides"]]},indent=2))


def build(args):
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen.canvas import Canvas

    data = read_json(args.data)
    assert data["model"]["bindings"]["model.onnx"] == MODEL_SHA
    assert len(data["slides"]) ==2 and sum(s["target_seconds"] for s in data["slides"]) ==120
    out = args.output
    out.mkdir(parents=True, exist_ok=True)
    fonts = args.font_dir
    regular = fonts/"arial.ttf"
    bold = fonts/"arialbd.ttf"
    if not regular.is_file() or not bold.is_file():
        raise FileNotFoundError("Supply --font-dir containing arial.ttf and arialbd.ttf.")
    image = Image.new("RGB", (1920,1080), BG)
    draw = ImageDraw.Draw(image)
    def text(value,x,y,size=32,color=INK,heavy=False,max_width=1760):
        font = ImageFont.truetype(str(bold if heavy else regular),size)
        bbox = draw.textbbox((0,0),value,font=font)
        width = bbox[2]-bbox[0]
        assert width <= max_width, (value,width,max_width)
        assert x+width <1900 and y+size <1065, value
        draw.text((x,y),value,font=font,fill=color)
    text("ORIGINAL EPOCH 44  /  FROZEN COMMAND-STAGE BENCHMARK",80,55,26,TEAL,True)
    text("Fast on PC. Human reliability is the gap.",80,115,68,INK,True)
    text(f"{data['original_recordings']:,} replayed  |  {data['scored_original_recordings']:,} scored  |  {data['ambiguous_original_recordings']:,} ambiguous / unscored  |  Previously exposed cohorts",80,208,29,MUTED)
    cols = (80,600,1050,1475)
    for title,x,width in zip(("Cohort","Correct accepted recall","Accepted precision","Negative false accepts"),cols,(480,410,385,365)):
        text(title,x,290,27,MUTED,True,width)
    draw.line((80,342,1840,342),fill="#30444E",width=2)
    for cohort,name,y in (("test/human","Human TEST",370),("holdout/human","Human HOLDOUT",505)):
        s=data["cohorts"][cohort]["selected"]
        text(name,cols[0],y+2,43,INK,True,480)
        text(f"{s['positives']:,} supported / {s['negatives']:,} negative",cols[0],y+64,25,MUTED,False,480)
        text(f"{s['recall']:.2%}",cols[1],y-4,62,AMBER if cohort.startswith('test') else TEAL,True,410)
        text(f"{s['precision']:.2%}",cols[2],y+4,49,INK,True,385)
        text(f"{s['negative_false_accepts']} / {s['negatives']}",cols[3],y+4,49,INK,True,365)
    synthetic = data["cohorts"]["test/synthetic"]["selected"]
    text(f"Synthetic TEST: {synthetic['recall']:.2%} recall (n={synthetic['positives']:,}); no negative evidence. Not proof of human reliability.",80,654,28,MUTED)
    for x in (80,680,1280):
        draw.rounded_rectangle((x,720,x+560,900),radius=18,fill="#14212A",outline="#30444E",width=2)
    latency = data["human_latency_ms"]
    text("Command stage — human recordings",105,741,27,MUTED,True,510)
    text(f"{latency['p50']:.2f} / {latency['p95']:.2f} ms",105,783,49,TEAL,True,510)
    text("p50 / p95; frontend + model + decision",105,851,23,MUTED,False,510)
    text("Isolated process peak",705,741,27,MUTED,True,510)
    text(f"{data['resource_probe']['peak_process_working_set_mib']:.2f} MiB",705,783,49,TEAL,True,510)
    text("Runtime + bounded preloaded probe audio",705,851,22,MUTED,False,510)
    text("Qualification",1305,741,27,MUTED,True,510)
    text("RESEARCH ONLY",1305,791,39,AMBER,True,510)
    text("No fresh blind-human / Pi / live proof",1305,851,24,MUTED,False,510)
    text("Intel i7-11700 CPU; 1 ONNX + 1 BLAS thread. Excludes wake inference, VAD, owner, UI and providers.",80,944,25,MUTED)
    text("Human-first, then human+synthetic training; DEV-selected epoch 44. Device actions stay disabled.",80,989,25,MUTED)

    poster = out/"slide-1.png"
    slide2 = out/"slide-2.png"
    with Image.open(args.poster) as img:
        assert img.size ==(1920,1080)
        img.convert("RGB").save(poster)
    image.save(slide2)
    movie = out/f"{PREFIX}-architecture.mp4"
    shutil.copyfile(args.video,movie)
    shutil.copyfile(args.data,out/"presentation-data.json")
    if args.gif:
        shutil.copyfile(args.gif,out/"architecture.gif")
    video = json.loads(subprocess.check_output([
        "ffprobe","-v","error","-select_streams","v:0","-show_entries",
        "stream=codec_name,width,height,avg_frame_rate:format=duration","-of","json",str(movie)
    ],text=True))
    stream=video["streams"][0]
    assert stream["codec_name"] =="h264" and (stream["width"],stream["height"]) ==(1920,1080)
    assert 15 < float(video["format"]["duration"]) <40

    prs=Presentation()
    prs.slide_width=Inches(40/3)
    prs.slide_height=Inches(7.5)
    prs.core_properties.title="OptionB Tiny epoch44 — two-minute presentation"
    prs.core_properties.subject="Two image-based slides with speaker notes and an embedded architecture movie"
    prs.core_properties.author="Smart Speaker project"
    prs.core_properties.keywords="epoch44,258912,ONNX,architecture,2-minute"
    for i,png in enumerate((poster,slide2)):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(png),0,0,width=prs.slide_width,height=prs.slide_height)
        if i ==0:
            slide.shapes.add_movie(str(movie),0,0,prs.slide_width,prs.slide_height,
                                   poster_frame_image=str(poster),mime_type="video/mp4")
        entry=data["slides"][i]
        note=(f"SLIDE{i+1} /2 — target {entry['target_seconds']}seconds\n"
              +( "0:00–1:05. Click the architecture to play. Static poster is sufficient if playback is unavailable.\n\n" if i==0 else "1:05–2:00. Do not describe exposed HOLDOUT as fresh blind evidence.\n\n")
              +entry["script"])
        slide.notes_slide.notes_text_frame.text=note
    pptx=out/f"{PREFIX}.pptx"
    prs.save(str(pptx))
    reread=Presentation(str(pptx))
    assert len(reread.slides) ==2
    for i,slide in enumerate(reread.slides):
        assert data["slides"][i]["script"] in slide.notes_slide.notes_text_frame.text
        for shape in slide.shapes:
            assert shape.left>=0 and shape.top>=0
            assert shape.left+shape.width <=reread.slide_width and shape.top+shape.height <=reread.slide_height
    with zipfile.ZipFile(pptx) as z:
        assert z.testzip() is None
        embedded=[n for n in z.namelist() if n.startswith("ppt/media/") and n.endswith(".mp4")]
        assert len(embedded)==1 and hashlib.sha256(z.read(embedded[0])).hexdigest()==sha(movie)
        assert len([n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml",n)])==2
        assert len([n for n in z.namelist() if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml",n)])==2

    pdf=out/f"{PREFIX}.pdf"
    c=Canvas(str(pdf),pagesize=(960,540))
    c.setTitle("OptionB Tiny epoch44 —2-slide static visual companion")
    c.setAuthor("Smart Speaker project")
    c.setSubject("Image-based visual companion generated from layout PNGs, not a native PowerPoint PDF export; no animation.")
    for png in (poster,slide2):
        c.drawImage(str(png),0,0,width=960,height=540)
        c.showPage()
    c.save()
    notes=out/f"{PREFIX}-speaker-notes.txt"
    notes.write_text("\n\n".join(
        f"SLIDE{s['number']}: target {s['target_seconds']}seconds\n\n{s['script']}"
        for s in data["slides"]
    )+"\n\nTiming is a suggested120-second speaking budget; rehearse. No narration recording is included.\n",encoding="utf-8")
    files=[poster,slide2,movie,pptx,pdf,notes,out/"presentation-data.json"]
    if args.gif:files.append(out/"architecture.gif")
    receipt={
        "format":"optionb-tiny-presentation-verification-v1", "slide_count":len(reread.slides),
        "suggested_speaking_seconds":sum(s["target_seconds"] for s in data["slides"]),
        "script_words":[s["script_word_count"] for s in data["slides"]],
        "video":video,"embedded_video_sha256_verified":True,"speaker_notes_roundtrip_verified":True,
        "native_powerpoint_render_performed":False,
        "visual_checks":"Inspect actual Manim frames, generated layout PNGs and rasterized companion PDF separately.",
        "pdf_note":"Static image-based visual companion, not a PowerPoint PDF conversion.",
        "model_source_commit":SOURCE_COMMIT,"model_sha256":MODEL_SHA,
        "files":{f.name:{"bytes":f.stat().st_size,"sha256":sha(f)} for f in files},
    }
    manifest=out/f"{PREFIX}-manifest.json"
    manifest.write_text(json.dumps(receipt,indent=2)+"\n",encoding="utf-8")
    archive=out/f"{PREFIX}.zip"
    with zipfile.ZipFile(archive,"w",zipfile.ZIP_DEFLATED,compresslevel=9) as z:
        for f in files+[manifest]:z.write(f,f.name)
    sums=out/f"{PREFIX}-SHA256SUMS"
    uploads=[pptx,pdf,movie,archive]
    sums.write_text("".join(f"{sha(f)}  {f.name}\n" for f in uploads),encoding="utf-8")
    uploads.append(sums)
    inventory={"repository":REPO,"release_tag":TAG,"model_source_commit":SOURCE_COMMIT,
               "assets":[{"name":f.name,"bytes":f.stat().st_size,"sha256":sha(f)} for f in uploads]}
    (out/"publication-asset-inventory.json").write_text(json.dumps(inventory,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"slide_count":2,"suggested_seconds":120,"words":receipt["script_words"],
                      "video":video,"assets":inventory["assets"]},indent=2))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    sub=parser.add_subparsers(dest="command",required=True)
    prep=sub.add_parser("prepare")
    prep.add_argument("--benchmark",type=Path,required=True)
    prep.add_argument("--output",type=Path,required=True)
    prep.add_argument("--frozen-source",type=Path,help="Pinned original snapshot checkout for the archived architecture and fit receipts")
    prep.set_defaults(func=prepare)
    render=sub.add_parser("build")
    render.add_argument("--data",type=Path,required=True)
    render.add_argument("--video",type=Path,required=True)
    render.add_argument("--poster",type=Path,required=True)
    render.add_argument("--gif",type=Path)
    render.add_argument("--output",type=Path,required=True)
    render.add_argument("--font-dir",type=Path,default=Path("C:/Windows/Fonts"))
    render.set_defaults(func=build)
    args=parser.parse_args()
    args.func(args)


if __name__=="__main__":
    main()
